"""Rebuild the review-item-2 draft from the archived original; outputs go only to 专项材料/第2条重建. Do not use this to overwrite a later current manuscript."""
from pathlib import Path
from copy import deepcopy
import re,json,hashlib
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

base=Path(__file__).resolve().parents[1]
out=base/'专项材料'/'第2条重建'
out.mkdir(parents=True,exist_ok=True)
original=base/'原始材料'/'频率分离双流网络的非接触式心率检测.docx'
raw=Document(original)
red='FF0000'
result=json.loads((base/'核验数据'/'paper_parameter_check.json').read_text())
slow,fast=result['results']
maxerr=max(v['max_absolute_amplitude_error'] for x in result['results'] for v in x['sinusoidal_checks'])

def mark(run):
 run.font.color.rgb=RGBColor.from_string(red)

def textp(doc,s,style=None):
 p=doc.add_paragraph(style=style)
 p.paragraph_format.space_after=Pt(5)
 p.paragraph_format.line_spacing=1.1
 p.paragraph_format.keep_together=True
 tokens=re.split(r'(Δts|xlf|xhf|H3D|Gσ|F0|Fs|νh|νw)',s)
 symbols={'Δts':('Δt','s'),'xlf':('x','lf'),'xhf':('x','hf'),'H3D':('H','3D'),'Gσ':('G','σ'),'F0':('F','0'),'Fs':('F','s'),'νh':('ν','h'),'νw':('ν','w')}
 for token in tokens:
  parts=[(token,False,False)] if token not in symbols else [(symbols[token][0],False,True),(symbols[token][1],True,token=='Gσ')]
  for content,issub,italic in parts:
   r=p.add_run(content);r.font.size=Pt(10.5);r.font.name='Times New Roman';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体');r.font.subscript=issub;r.font.italic=italic;mark(r)
 return p

def mr(s,plain=False):
 r=OxmlElement('m:r')
 if plain:
  prop=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');prop.append(sty);r.append(prop)
 rp=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),red);rp.append(color)
 font=OxmlElement('w:rFonts');font.set(qn('w:ascii'),'Cambria Math');font.set(qn('w:hAnsi'),'Cambria Math');rp.append(font)
 size=OxmlElement('w:sz');size.set(qn('w:val'),'21');rp.append(size);r.append(rp)
 t=OxmlElement('m:t');t.text=s;r.append(t);return r

def pack(tag,children):
 el=OxmlElement('m:'+tag)
 if not isinstance(children,list):children=[children]
 for x in children:el.append(x)
 return el

def sub(base,index):return pack('sSub',[pack('e',mr(base)),pack('sub',mr(index,plain=index in {'lf','hf','3D','cascade','t','h','w'}))])
def sup(base,power):return pack('sSup',[pack('e',mr(base)),pack('sup',mr(power))])
def frac(a,b):return pack('f',[pack('num',a),pack('den',b)])
def exp(expr):return [mr('exp',True),mr('('),*expr,mr(')')]
def sumexpr(lower,upper,expr,op='∑'):
 pr=OxmlElement('m:naryPr');c=OxmlElement('m:chr');c.set(qn('m:val'),op);pr.append(c)
 loc=OxmlElement('m:limLoc');loc.set(qn('m:val'),'subSup');pr.append(loc)
 return pack('nary',[pr,pack('sub',mr(lower)),pack('sup',mr(upper)),pack('e',expr)])

def eq(doc,num,parts):
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.LEFT
 p.paragraph_format.tab_stops.add_tab_stop(Cm(8.3),WD_TAB_ALIGNMENT.CENTER)
 p.paragraph_format.tab_stops.add_tab_stop(Cm(16.6),WD_TAB_ALIGNMENT.RIGHT)
 p.add_run('\t')
 p.paragraph_format.space_after=Pt(7);p.paragraph_format.space_before=Pt(4)
 rp=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),red);rp.append(c);p._p.get_or_add_pPr().append(rp)
 om=OxmlElement('m:oMath')
 for x in parts:om.append(x)
 p._p.append(om);r=p.add_run(f'\t({num})');mark(r);r.font.size=Pt(10.5)
 return p

def setup(doc):
 sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=sec.right_margin=Cm(2)
 st=doc.styles['Normal'];st.font.name='Times New Roman';st.font.size=Pt(10.5);st._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
 for name in ['Title','Heading 1','Heading 2']:
  st=doc.styles[name];st.font.color.rgb=RGBColor(0,0,0);st._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体')

# Standalone exact replacement section; shared with full manuscript.
section=Document();setup(section)
h=textp(section,'2.2.1　时空频率分解与离散响应');h.runs[0].bold=True;h.paragraph_format.keep_with_next=True
textp(section,'给定平滑尺度σ，取奇数核长K=2⌈2σ⌉+1，令M=(K−1)/2。以零为中心的离散坐标n取−M至M，归一化一维高斯核如式(1)所示；其系数和为1，“零中心”不表示核系数的均值为零。以一维核的外积构造三维核Gσ，如式(2)所示。')
num=exp([frac([mr('−'),sup('n','2')],[mr('2'),sup('σ','2')])])
den=sumexpr('k=−M','M',exp([frac([mr('−'),sup('k','2')],[mr('2'),sup('σ','2')])]))
eq(section,1,[sub('g','σ'),mr('[n] = '),frac(num,[den]),mr(',  −M ≤ n ≤ M')])
eq(section,2,[sub('G','σ'),mr('[i,j,k] = '),sub('g','σ'),mr('[i]'),sub('g','σ'),mr('[j]'),sub('g','σ'),mr('[k]')])
textp(section,'三维核通过逐通道卷积作用于输入特征x，得到平滑分量xlf及残差分量xhf，如式(3)所示。其中∗表示三维卷积；同一离散尺度作用于时空三轴，但时间轴与空间轴的物理单位不同，后续时间频率换算使用特征序列的采样间隔。')
eq(section,3,[sub('x','lf'),mr(' = '),sub('G','σ'),mr(' ∗ x,    '),sub('x','hf'),mr(' = x − '),sub('x','lf')])
textp(section,'设视频原始帧率为F0，分支的时间压缩倍数为s，则名义特征采样间隔Δts和采样率Fs由式(4)给出。F0=30 Hz时，慢流s=4、Fs=7.5 Hz，快流s=2、Fs=15 Hz。采样率换算仅确定频率坐标，不能据此认定前级降采样已消除混叠。')
eq(section,4,[sub('Δt','s'),mr(' = '),frac([mr('s')],[sub('F','0')]),mr(',    '),sub('F','s'),mr(' = '),frac([sub('F','0')],[mr('s')])])
textp(section,'在忽略有限特征图填充边界的条件下，一维离散核的时间频率响应为式(5)，f的单位为Hz，j为虚数单位。响应由σ、核长K和Fs共同决定，不能只由σ给出物理频率。')
eq(section,5,[sub('H','σ,K'),mr('(f;'),sub('F','s'),mr(') = '),sumexpr('n=−M','M',[sub('g','σ'),mr('[n]'),*exp([frac([mr('−j2πfn')],[sub('F','s')])])])])
textp(section,'三维核的响应可分离为式(6)，其中νh、νw为空间频率，单位为周期/特征像素。空间频率为零时，归一化空间核的响应为1，式(5)即对应三维算子的时间轴截面；该截面不代表所有时空输入的响应。残差算子的三维响应为1−H3D。')
eq(section,6,[sub('H','3D'),mr('(f,'),sub('ν','h'),mr(','),sub('ν','w'),mr(') = '),sub('H','t'),mr('(f)'),sub('H','h'),mr('('),sub('ν','h'),mr(')'),sub('H','w'),mr('('),sub('ν','w'),mr(')')])
textp(section,'本方法慢流σ=1.8、K=9，快流σ=0.8、K=5。定义幅值降至50%的频率满足|H(f)|=0.5，功率降至50%的频率满足|H(f)|²=0.5，分别对应约−6.02 dB与−3.01 dB。由实际有限长离散核计算，慢流的两个频率分别为0.799 Hz和0.571 Hz，快流分别为3.520 Hz和2.493 Hz。因此0.5 Hz不作为这两组参数的共同截止频率。')
textp(section,'未截断解析高斯的半高全宽为2√(2ln2)σ，慢流与快流分别约为4.239和1.884个特征采样点。按各分支采样间隔换算，分别约为565 ms和126 ms。该宽度用于解释尺度参数的时间量纲；有限截断、离散归一化后的截止特性应以式(5)计算，半高全宽不能直接替代截止频率。')
textp(section,'多层重复作用需区分纯滤波串联与实际网络计算。若m个相同线性滤波器在同一采样率下直接串联，则其响应满足式(7)；在非相同滤波器的情形，应使用各层响应的乘积。图2给出m=1、2、3、6时的纯滤波参考曲线，用于说明重复平滑会进一步改变幅频特性。')
eq(section,7,[sub('H','cascade'),mr('(f) = '),sumexpr('ℓ=1','m',[sub('H','ℓ'),mr('(f)')],op='∏'),mr(' = '),pack('sSup',[pack('e',[mr('['),sub('H','σ,K'),mr('(f;'),sub('F','s'),mr(')]')]),pack('sup',mr('m'))])])
textp(section,'实际慢流与快流各包含3个FreqSepBlock；每个块除分解滤波外，低频路径还有一次同尺度高斯平滑。从低频处理路径计，每条流包含6处高斯滤波，但这些算子之间穿插卷积、归一化、激活、BiMamba、通道门控及跨分支交换，块输出还要经过分支求和和跨流连接。因此不能用式(7)的m=6结果代表整个分支或完整网络，也不能把3个FreqSepBlock等同于3次纯低通串联。上采样与输出阶段的平滑位于其他特征尺度，其影响同样不并入单个分解核的截止频率。')
textp(section,'离散响应及算子数值核验如图2所示。核验采用空间恒定的正弦输入，频率覆盖两条流的主要分析区间；每个序列含128个特征采样点，使用单精度浮点运算。对连续1、2、3次纯滤波分别剔除时空边界影响后拟合正弦幅值，在每条流12个频率点、共72组幅值比较中，计算曲线与滤波输出的最大绝对幅值误差为1.26×10⁻⁶。该核验检验有限离散核及其纯串联计算的一致性，不测量训练后完整网络的频率增益。')
p=section.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(base/'图表'/'图2_按现稿参数的离散频率响应.png'),width=Cm(16.8))
p=textp(section,'图2　按稿件参数计算的高斯核离散幅频响应与纯滤波串联参考');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
textp(section,'注：(a)(b)分别为慢流与快流的单层低通、残差响应，圆点为滤波输出的正弦幅值核验；(c)(d)为相同采样率下1、2、3、6次纯线性滤波串联参考，其中1、2、3次含数值核验点。所有曲线取空间零频率并忽略边界，6次曲线不表示完整网络响应。')
section.save(out/'第2条_替换章节_标红.docx')

# Full manuscript: retain original experiments, configurations, citations and existing tables.
d=Document(original)
oldparas=list(d.paragraphs)
changes=[]
# Renumber original equations 4--9 to 8--13, preserving MathType objects.
for t in [t for p in d.paragraphs[55:65] for t in p._p.xpath('.//w:t')]:
 old=t.text or '';new=re.sub(r'\(([4-9])\)',lambda m:f'({int(m.group(1))+4})',old)
 if new!=old:
  t.text=new
  r=t.getparent()
  if r.tag==qn('w:r'):
   rp=r.find(qn('w:rPr'))
   if rp is None:rp=OxmlElement('w:rPr');r.insert(0,rp)
   c=rp.find(qn('w:color'))
   if c is None:c=OxmlElement('w:color');rp.append(c)
   c.set(qn('w:val'),red)

def replace_text(p,s):
 p.clear();r=p.add_run(s);mark(r)

edits={
5:[('截止频率由生理频带锚定，不随训练域漂移','滤波参数跨数据集固定，时间幅频响应由核长、平滑尺度与特征采样率共同决定')],
}
# Match by content rather than paragraph indices for coordinated clause edits.
for p in d.paragraphs:
 s=p.text;new=s
 new=new.replace('截止频率由生理频带锚定，不随训练域漂移','滤波参数跨数据集固定，时间幅频响应由核长、平滑尺度与特征采样率共同决定')
 new=new.replace('with the cut-off frequency anchored to the physiological heart-rate band and thus invariant to training-domain noise.','with filter parameters held fixed across datasets; the discrete temporal response is jointly determined by kernel length, Gaussian scale, and feature sampling rate.')
 new=new.replace('截止频率直接由0.5～3 Hz心搏生理频带给出','时间幅频响应由核长、平滑尺度与分支采样率共同确定')
 new=new.replace('，且不施加额外平滑以免损失波形细节。','；快流内部的高频路径不施加额外平滑，低频路径仍采用σ=0.8的高斯平滑。')
 new=new.replace('平滑仅消除上采样引入的阶梯伪影，对0.5～3 Hz心搏主频衰减较小；','平滑对不同时间频率具有不同衰减，其作用不能限定为仅消除上采样伪影，分解核的离散响应分析见第2.2.1节；')
 if s.startswith('频率分离模块FreqSepBlock是本文方法的核心组件'):
  new='频率分离模块FreqSepBlock是本文方法的核心组件，在时序编码前采用固定三维高斯核进行平滑与残差分解。固定核权重不参与梯度更新；不同分支的时间响应由核长、平滑尺度及特征采样率共同决定，需按各自时间尺度进行离散频率分析。'
 new=new.replace('与高斯平滑层(σ=1.8)','与高斯平滑层（慢流σ=1.8，快流σ=0.8）')
 new=new.replace('一维高斯平滑(σ=1.5)仅作用于上采样阶梯伪影，在0.5～3 Hz心搏主频处响应基本不变，故该差距可直接归因于阶梯伪影。','一维高斯平滑(σ=1.5)具有频率依赖的衰减，因此该消融结果支持平滑回归头对预测误差的影响，但不能把差距唯一归因于上采样阶梯伪影，也不能据此认为0.5～3 Hz范围内的响应保持不变。')
 new=new.replace('本方法以0.5～3 Hz心搏频带给出的单一固定边界换取跨域稳定性','本方法采用跨数据集固定尺度的高斯分解核')
 if new!=s:
  changes.append({'before':s,'after':new});replace_text(p,new)

# Replace original subsection 2.2.1, including its three formula paragraphs.
start=next(p for p in d.paragraphs if p.text.startswith('2.2.1'))
end=next(p for p in d.paragraphs if p.text.startswith('2.2.2'))
startel=start._p;endel=end._p;parent=startel.getparent()
idx=parent.index(startel)
while parent[idx] is not endel:parent.remove(parent[idx])
# Place the longer analysis in a single-column section; retain original two-column layout elsewhere.
section_template=deepcopy(d.sections[-1]._sectPr)
def sectionbreak(cols):
 p=OxmlElement('w:p');pr=OxmlElement('w:pPr');sec=deepcopy(section_template)
 for node in list(sec):
  if node.tag in [qn('w:headerReference'),qn('w:footerReference'),qn('w:titlePg')]:sec.remove(node)
 typ=sec.find(qn('w:type'))
 if typ is None:typ=OxmlElement('w:type');sec.insert(0,typ)
 typ.set(qn('w:val'),'continuous')
 cs=sec.find(qn('w:cols'))
 if cs is None:cs=OxmlElement('w:cols');sec.append(cs)
 cs.set(qn('w:num'),str(cols))
 pr.append(sec);p.append(pr);return p
endel.addprevious(sectionbreak(2))
# Copy paragraphs, then replace image relationship with an image part in the target package.
for element in section.element.body:
 if element.tag==qn('w:sectPr'):continue
 clone=deepcopy(element)
 for blip in clone.xpath('.//a:blip'):
  rid,_=d.part.get_or_add_image(str(base/'图表'/'图2_按现稿参数的离散频率响应.png'))
  blip.set(qn('r:embed'),rid)
 endel.addprevious(clone)
endel.addprevious(sectionbreak(1))
# Disable tracked-changes mode; red direct formatting is used.
for node in list(d.settings.element):
 if node.tag==qn('w:trackRevisions'):d.settings.element.remove(node)
d.save(out/'频率分离双流网络的非接触式心率检测_第2条修订_标红.docx')
(out/'coordinated_changes.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2))

# Reviewer response, based on actual completed calculations only.
reply=Document();setup(reply)
reply.add_heading('修改说明：审回意见第2条',0)
reply.add_paragraph('稿件编号：9953480\n论文题目：频率分离双流网络的非接触式心率检测\n本文件为本轮第2条意见的专项修改说明，后续可并入完整逐条回复。')
reply.add_heading('审稿意见',1)
reply.add_paragraph('第2.2节称σ=1.8对应4.2帧、140 ms，并在约0.52 Hz处达到50%衰减，该换算关系需严格核验。请给出考虑帧率、时间步长和多层重复滤波后的离散频率响应推导，并绘制实际幅频曲线，修正不一致的数值。')
reply.add_heading('回复与修改依据',1)
reply.add_paragraph('感谢审稿专家指出时间尺度和频率响应说明中的问题。原稿将高斯尺度对应的特征采样点直接按原始视频帧率换算为时间，未充分区分时间压缩后的特征采样率；同时未明确“50%衰减”的幅值或功率定义，所列频率数值与稿件规定的有限离散核不一致。本次按现稿的实验结果及方法参数进行核验，保留σ=1.8/0.8、K=2⌈2σ⌉+1、原始帧率30 Hz和慢/快流4倍/2倍时间压缩，重新推导和计算频率响应。')
reply.add_heading('1．补充有限离散核及采样间隔推导',2)
reply.add_paragraph('修订第2.2.1节，以式(1)—式(3)明确归一化离散核、三维外积及分解方式；式(4)给出Δts=s/F0及Fs=F0/s；式(5)给出有限长离散核的频率响应，式(6)解释三维响应与空间零频率时间截面的关系。慢流K=9、Fs=7.5 Hz，快流K=5、Fs=15 Hz。补充说明：时间响应由核长、σ和分支采样率共同决定，空间与时间频率的物理单位不同。')
reply.add_heading('2．修正半高全宽及50%衰减的数值',2)
t=reply.add_table(rows=1,cols=5);t.style='Table Grid'
for c,s in zip(t.rows[0].cells,['分支','核长/采样率','解析半高全宽换算','幅值50%频率','功率50%频率']):c.text=s
for x in result['results']:
 cells=t.add_row().cells
 for c,s in zip(cells,[x['stream'],f"K={x['kernel']} / {x['feature_fs_Hz']:g} Hz",f"{x['untruncated_FWHM_feature_samples']:.3f}点 / {x['untruncated_FWHM_ms']:.0f} ms",f"{x['half_amplitude_Hz']:.3f} Hz",f"{x['half_power_Hz']:.3f} Hz"]):c.text=s
reply.add_paragraph('原稿中的慢流约140 ms、快流约63 ms分别更正为约565 ms和126 ms，并将“帧”明确为“特征采样点”。上述半高全宽来自未截断解析高斯，仅用于尺度量纲说明；有限离散核的截止特性以离散响应计算。原0.52 Hz和1.17 Hz不再作为两组核的50%衰减频率。幅值50%与功率50%分别定义为|H|=0.5和|H|²=0.5，对应约−6.02 dB和−3.01 dB。')
reply.add_heading('3．说明多层重复滤波的适用条件',2)
reply.add_paragraph('新增式(7)，说明相同采样率下的纯线性滤波器串联响应等于各层响应的乘积，相同核时为H的m次方。新增图2展示1、2、3、6次纯滤波的重复作用。每个FreqSepBlock包含分解滤波和低频路平滑，每条流3个块沿低频处理路径包含6处高斯滤波；但中间存在非线性编码、分支求和和跨流交互，故不将H³或H⁶表述为完整网络响应。上采样及输出阶段的平滑也不归入单个分解核的截止频率。')
reply.add_heading('4．绘制响应曲线并完成滤波算子数值核验',2)
reply.add_paragraph(f'图2(a)(b)使用有限离散核系数直接计算慢、快流低通及残差响应；图2(c)(d)给出纯线性串联参考曲线。另以空间恒定的合成正弦输入进行算子级数值核验，每条流12个频率点，每个序列128个特征采样点，比较单次及连续2、3次滤波，共72组幅值结果；使用CPU单精度计算，并剔除重复滤波的时空边界影响，以正弦/余弦及常数项最小二乘拟合幅值。理论曲线与数值输出的最大绝对幅值误差为{maxerr:.3g}，支持计算与滤波算子的一致性。')
reply.add_paragraph('验证对象为按稿件参数显式配置的高斯滤波算子及纯线性串联。该结果不是训练后完整网络的经验传递函数，也不替代心搏频带能量或运动泄漏分析；后者将在第3条意见中处理。完整数值结果及生成脚本随本轮材料提供。')
reply.add_heading('5．同步修正文中关联表述',2)
reply.add_paragraph('同步调整中英文摘要、第1.5节、第2.1节、第2.2节及第3.7.2节中将截止频率仅归因于σ或直接由0.5—3 Hz给出的表述；补充第2.2.2节中慢/快流低频路径的不同平滑尺度。第2.1节及第3.4节删除“平滑仅作用于上采样伪影”“心搏频带响应基本不变”的确定性解释，保留原消融数值，并限定其支持的结论范围。')
reply.add_heading('修改位置与核对',1)
reply.add_paragraph('主要修改位于修订稿第2.2.1节、式(1)—式(7)及新增图2；关联位置见上文第5点。原式(4)—式(9)顺延为式(8)—式(13)，对应文字引用同步更新。全部新写和修改内容直接标红，未使用修订模式。原表1—表9及其中数据保持不变，训练配置和实验指标未因本次频率换算核验而重新估计。')
reply.save(out/'修改说明_第2条.docx')

# Artifact checks.
new=Document(out/'频率分离双流网络的非接触式心率检测_第2条修订_标红.docx')
assert [p.text for p in raw.paragraphs if re.match(r'^\[\d+\]',p.text)] == [p.text for p in new.paragraphs if re.match(r'^\[\d+\]',p.text)], 'Reference entries changed'
assert len(raw.tables)==len(new.tables)==9
assert [t._tbl.xml for t in raw.tables]==[t._tbl.xml for t in new.tables], 'Original tables changed'
assert len(new.element.body.xpath('.//m:oMath'))==7
assert len(new.element.body.xpath('.//w:ins | .//w:del'))==0
checks={'original_sha256':hashlib.sha256(original.read_bytes()).hexdigest(),'original_table_count':9,'original_table_xml_unchanged':True,'new_editable_equations':7,'tracked_revisions':0,'coordinated_paragraph_changes':len(changes),'basis':'Original manuscript experimental results and method parameters, confirmed by user; frequency/time conversions recomputed.'}
(out/'document_checks.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2))
print(json.dumps(checks,ensure_ascii=False))
